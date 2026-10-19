#!/usr/bin/env python3
"""Backup (補足) slides for the 2026-10-18 midterm talk, in the design of the user's deck.

Slides: per-sequence HOTA difference (native chart), evaluation conditions, mask-candidate score and its weight,
constant-velocity example (KF), Mamba block of the motion predictor (copied from the restructured deck),
method-2 pilot on MOSE, checkpoint comparison. Numbers come from TrackEval outputs / experiment logs (see manifest.md).
usage: make_backup_slides.py   (python-pptx required)
"""
import copy
import csv
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
MAT = HERE.parents[1]
PROJ = MAT.parent
USER_DECK = MAT / "20261018_中間発表_油谷.pptx"         # デザイン（テーマ・レイアウト）だけ借りる
RESTRUCTURED = MAT / "2026-10-12-midterm-slides-restructured.pptx"
OUT = HERE / "2026-10-18-midterm-backup-slides.pptx"
TE = Path("/mnt/HDD10TB-2/aburatani/TrackEval/data/trackers/dancetrack/val")
OURS_DETAILED = Path("/mnt/HDD10TB-2/aburatani/2025_09_aburatani_Mamba_Trackers/artifacts/sam2_minimal_integration/trackeval/"
                     "sam2_stateful_minimal_20260723T_sam2_minimal_25seq_epoch100_seqinfo/"
                     "sam2_stateful_minimal_20260723T_sam2_minimal_25seq_epoch100/pedestrian_detailed.csv")
KF_FIG = MAT / "figures/2026-10-14-kf-constant-velocity-example/kf-constant-velocity.png"
BLUE, RED, MUTED, TXT = RGBColor(0x2E, 0x75, 0xB6), RGBColor(0xC0, 0x00, 0x00), RGBColor(0x59, 0x59, 0x59), RGBColor(0x26, 0x26, 0x26)

prs = Presentation(str(USER_DECK))
lst = prs.slides._sldIdLst
for el in list(lst):  # スライドは全部外し，テーマとレイアウトだけ使う
    lst.remove(el); prs.part.drop_rel(el.get(qn("r:id")))
LAYOUT = {l.name: l for l in prs.slide_layouts}


def new_slide(title, layout="タイトルのみ"):
    s = prs.slides.add_slide(LAYOUT[layout])
    s.shapes.title.text_frame.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(40)
    return s


def body(s, items, top=1.2, height=None, left=0.36, width=12.66):
    """本文（レベル0 は28pt，レベル1 は24pt）"""
    ph = next(p for p in s.placeholders if p.placeholder_format.idx)
    ph.left, ph.top, ph.width = Inches(left), Inches(top), Inches(width)
    ph.height = Inches(height or 0.55 * len(items))
    tf = ph.text_frame; tx = tf._txBody
    for na in tx.iter(qn("a:normAutofit")):
        for a in ("fontScale", "lnSpcReduction"):
            na.attrib.pop(a, None)
    proto = copy.deepcopy(tx.findall(qn("a:p"))[0])
    for ch in list(proto):
        if ch.tag != qn("a:pPr"):
            proto.remove(ch)
    for p in tx.findall(qn("a:p")):
        tx.remove(p)
    for lvl, text in items:
        el = copy.deepcopy(proto); tx.append(el)
        p = _Paragraph(el, tf); p.level = lvl
        r = p.add_run(); r.text = text; r.font.size = Pt(28 if lvl == 0 else 24)


def note(s, x, y, w, h, text, size=14, color=MUTED, bold=False, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    return tb


def table(s, rows, x, y, w, col_w=None, size=20, bold_cells=()):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.45 * len(rows))).table
    if col_w:
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = str(v)
            p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(size); r.font.bold = (i == 0) or ((i, j) in bold_cells)
    return t


def seq_hota(path):
    return {r["seq"]: float(r["HOTA___AUC"]) * 100 for r in csv.DictReader(open(path)) if r["seq"] != "COMBINED"}


# ---- 1) 系列ごとの HOTA の差
samurai = seq_hota(TE / "my_samurai_model/pedestrian_detailed.csv")
ours = seq_hota(OURS_DETAILED)
diff = sorted(((ours[k] - samurai[k], k) for k in ours), reverse=True)
s = new_slide("補足：系列ごとのHOTAの差")
cd = CategoryChartData()
cd.categories = [k.replace("dancetrack", "") for _, k in diff]
cd.add_series("HOTA の差", [round(d, 2) for d, _ in diff])
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.95), Inches(12.3), Inches(4.9), cd).chart
ch.has_legend = False
ch.has_title = False
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW  # 負の棒に番号が重ならないよう下端に
plot = ch.plots[0]; plot.gap_width = 40
for i, (d, _) in enumerate(diff):
    pt = plot.series[0].points[i]
    pt.format.fill.solid(); pt.format.fill.fore_color.rgb = BLUE if d >= 0 else RED
ch.category_axis.tick_labels.font.size = Pt(12)
ch.value_axis.tick_labels.font.size = Pt(14)
ch.value_axis.has_major_gridlines = True
ch.category_axis.has_title = True; ch.category_axis.axis_title.text_frame.text = "系列（dancetrack の番号）"
ch.value_axis.has_title = True; ch.value_axis.axis_title.text_frame.text = "HOTA の差（Ours − SAMURAI）"
for ax in (ch.category_axis, ch.value_axis):
    for r in ax.axis_title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(14); r.font.bold = False
win = sum(d > 0 for d, _ in diff)
tiny = sum(abs(d) < 0.1 for d, _ in diff)
note(s, 0.5, 1.15, 12.3, 0.8,
     f"DanceTrack val 25系列のうち {win} 系列で Ours（state）が SAMURAI を上回る（青：上回る，赤：下回る。差が 0.1 未満の系列が {tiny} つ）\n"
     f"差は {diff[-1][0]:+.1f}（{diff[-1][1]}）〜 {diff[0][0]:+.1f}（{diff[0][1]}）。25系列まとめた HOTA は 54.11 → 55.52（+1.41）",
     size=16, color=TXT)
note(s, 0.5, 6.95, 12.3, 0.3, "系列ごとの HOTA は TrackEval の結果（定量的評価の表と同じ推論結果）。1回の評価で，ばらつきは未確認", size=11)

# ---- 2) 評価の条件
s = new_slide("補足：評価の条件", "タイトルとコンテンツ")
body(s, [(0, "データ：DanceTrack val 25系列（2〜19人のダンス）"),
         (0, "初期化：全対象を，最初に出てくるフレームで正解の bbox で指定"),
         (1, "検出器から始める一般的な MOT の評価とは条件が違う（OC-SORT などと直接は比べられない）"),
         (0, "比較手法も同じ条件"),
         (1, "SAMURAI（SOT）は物体ごとに並列に動かす（物体ごとにメモリとカルマンフィルタ）"),
         (1, "提案手法も同じ並列化。違いは動き予測（カルマンフィルタ → Mamba）だけ"),
         (0, "モデル：SAM2 は tiny（追加学習なし）"),
         (1, "Mamba は DanceTrack の学習用データの正解 bbox 系列で単体学習（100 epoch）"),
         (0, "評価：TrackEval（HOTA / DetA / AssA / IDF1 / IDSW）")], top=1.2, height=5.9)

# ---- 3) マスク候補選択のスコアと重み
s = new_slide("補足：マスク候補選択のスコアと重み", "タイトルとコンテンツ")
body(s, [(0, "SAM2 が出す3つのマスク候補から，スコア最大のものを選ぶ"),
         (1, "score ＝ w × motion score ＋（1 − w）× affinity score"),
         (1, "motion score：Mamba の予測 bbox と候補の bbox の IoU"),
         (1, "affinity score：SAM2 自身が出す候補の信頼度（予測 IoU）"),
         (1, "SAMURAI と同じ式・同じ重み（w ＝ 0.25）。違いは予測を出す部分だけ")], top=1.2, height=2.6)
rows = [["重み w", "HOTA↑", "DetA↑", "AssA↑", "IDF1↑", "IDSW↓"],
        ["0.25（表の Ours）", "55.52", "49.60", "62.48", "64.15", "1535"],
        ["0.5", "53.92", "48.35", "60.45", "61.73", "1559"],
        ["参考：SAMURAI（KF, w ＝ 0.25）", "54.11", "47.89", "61.48", "62.33", "1517"]]
table(s, rows, 0.6, 4.05, 12.1, col_w=[4.1, 1.6, 1.6, 1.6, 1.6, 1.6], size=18,
      bold_cells=[(1, j) for j in range(1, 5)])
note(s, 0.6, 5.95, 12.1, 0.9,
     "Ours（state，100 epoch）で w だけを変えた結果（同じコード）。w を 0.5 にすると SAMURAI を下回る\n"
     "→ 重みで結果が大きく変わる。w ＝ 0（予測を使わない）との比較はまだない（重みの決め方とあわせて今後の課題）", size=14, color=TXT)

# ---- 4) 等速の仮定がずれる例
s = new_slide("補足：等速の仮定による予測のずれ")
pic = s.shapes.add_picture(str(KF_FIG), Inches(0.9), Inches(1.6))
pic.width = Inches(11.5); pic.height = int(Inches(11.5) * 600 / 1624)
y = 1.6 + 11.5 * 600 / 1624 + 0.2
note(s, 0.9, y, 11.5, 1.6,
     "左：時刻 t。白：これまでの軌跡，赤の点線：これまでの平均速度のまま進むと仮定した予測\n"
     "右：0.75 秒後（15 フレーム後）。緑：実際の位置，赤の点線：等速を仮定した予測（約 370 px ずれる）\n"
     "ダンスでは急に止まる・向きを変える動きが多く，等速（カルマンフィルタの運動モデル）の仮定がずれやすい",
     size=16, color=TXT)
note(s, 0.9, 7.0, 11.5, 0.3, "dancetrack0052（train）の正解データ（ID 3，frame 264）から作図した模式的な例で，SAMURAI の出力ではない", size=11)

# ---- 5) 動き予測器の Mamba ブロック（再構成版から複写）
src_prs = Presentation(str(RESTRUCTURED))
src = next(x for x in src_prs.slides if x.shapes.title is not None and x.shapes.title.text_frame.text == "補足：動き予測器のMambaブロック")
s = new_slide("補足：動き予測器のMambaブロック")
for el in src.shapes._spTree:
    if el.tag not in (qn("p:sp"), qn("p:pic"), qn("p:cxnSp"), qn("p:grpSp")) or el.find(".//" + qn("p:ph")) is not None:
        continue
    el = copy.deepcopy(el)
    for node in el.iter():
        for k, v in list(node.attrib.items()):
            if k in (qn("r:embed"), qn("r:link"), qn("r:id")) and v:
                rel = src.part.rels[v]
                if rel.reltype.endswith("/image"):
                    _, rid = s.part.get_or_add_image_part(__import__("io").BytesIO(rel.target_part.blob))
                    node.set(k, rid)
                else:
                    node.set(k, s.part.relate_to(rel.target_part, rel.reltype))
    s.shapes._spTree.append(el)
note(s, 9.3, 2.0, 3.8, 2.4, "構成（state carry 型）\nd_model 256，d_state 16\n入力：対象ごとの bbox\n出力：次フレームでの bbox の変化 Δb̂", size=16, color=TXT)

# ---- 5b) 手法1の Mamba の学習方法（state carry 型，表の Ours（state）の 100 epoch のモデル）
s = new_slide("補足：手法1のMambaの学習方法", "タイトルとコンテンツ")
body(s, [(0, "データ：DanceTrack 学習用の正解bbox"),
         (1, "417軌跡（約35万フレーム），評価のvalとは別"),
         (0, "入力：12フレームの bbox 系列"),
         (1, "画像サイズで正規化した (x, y, w, h)"),
         (0, "出力：次のフレームとの差分 Δb"),
         (1, "Smooth L1 損失（最初の4ステップは除く）"),
         (0, "窓ごとに状態を初期化して学習"),
         (1, "推論では，状態を系列全体で引き継ぐ")], top=1.2, height=5.6, width=7.7)
rows = [["設定", "値"], ["Mamba の層数", "3"], ["d_model / d_state", "256 / 16"], ["最適化", "Adam"],
        ["学習率", "1e-4"], ["バッチサイズ", "64"], ["epoch", "100"], ["勾配クリップ", "1.0"], ["教師の Δb の倍率", "50"]]
table(s, rows, 8.4, 1.45, 4.5, col_w=[2.6, 1.9], size=16)
note(s, 8.4, 5.7, 4.5, 1.2, "推論では学習と違い，窓で区切らずに状態を引き継ぐ（学習と推論の条件が違う点は，今後の学習方法の見直しで扱う）", size=13)
note(s, 0.5, 7.0, 12.3, 0.3, "設定は学習時に保存された MambaStateful.yaml（saved_ckpts/mamba_stateful_dancetrack）と学習コード（StatefulTrajDataset）から", size=11)

# ---- 5c) 手法1の Mamba の入出力（推論時，1つの物体・時刻 t）。表の Ours（state）の実行時の設定・コードに合わせる
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from lxml import etree
import re as _re
s = new_slide("補足：手法1のMambaの入出力（1つの物体）")
SAM_F, SAM_L = RGBColor(0xF8, 0xCB, 0xAD), RGBColor(0xED, 0x7D, 0x31)
MB_F, MB_L = RGBColor(0xDE, 0xEB, 0xF7), RGBColor(0xC0, 0x00, 0x00)
SEL_F, SEL_L = RGBColor(0xE2, 0xF0, 0xD9), RGBColor(0x70, 0xAD, 0x47)
GR_F, GR_L = RGBColor(0xF2, 0xF2, 0xF2), RGBColor(0x8C, 0x8C, 0x8C)


def rich(p, text, size, bold, color=TXT):
    """「b_{t-1}^{k}」のような書き方を，下付き・上付きの文字にして段落へ入れる"""
    for m in _re.finditer(r"([_^])\{([^}]*)\}|([^_^]+|[_^])", text):  # 波かっこ付きの _{…}・^{…} だけを下付き・上付きにする
        if m.group(3) is not None:
            seg, base = m.group(3), None
        else:
            seg = m.group(2); base = "-25000" if m.group(1) == "_" else "30000"
        r = p.add_run(); r.text = seg; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        if base:
            r._r.get_or_add_rPr().set("baseline", base)


def box(x, y, w, h, lines, fill, line, size=15, lw=1.5):
    bx = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    bx.adjustments[0] = 0.12
    bx.fill.solid(); bx.fill.fore_color.rgb = fill
    bx.line.color.rgb = line; bx.line.width = Pt(lw); bx.shadow.inherit = False
    tf = bx.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, ln in enumerate(lines):
        text, bold = (ln if isinstance(ln, tuple) else (ln, False))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
        rich(p, text, size, bold)
    return bx


def arrow(pts, head=True, color=RGBColor(0x40, 0x40, 0x40), w=1.75):
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = color; c.line.width = Pt(w)
        if head and i == len(pts) - 2:
            e = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd")); e.set("type", "triangle")


Y1, Y2, Y3, Y4 = 1.55, 3.0, 4.4, 5.65   # 各段の上端
# 1段目：SAM2 側
box(0.35, Y1, 2.0, 0.9, [("SAM2", True), "mask decoder"], SAM_F, SAM_L)
box(2.75, Y1, 2.8, 0.9, [("マスク候補 ×3", True), "＋ 信頼度 s_{mask}^{k}（k ＝ 1, 2, 3）"], SAM_F, SAM_L, size=14)
box(5.95, Y1, 2.4, 0.9, [("候補の bbox b_{t}^{k}", True), "（マスクの外接矩形）"], SAM_F, SAM_L)
# 2段目：Mamba 側
box(0.35, Y2, 2.0, 0.9, [("Mamba の状態", True), "h_{t−1}（物体ごと）"], MB_F, MB_L, size=14, lw=2.5)
box(2.75, Y2, 2.8, 0.9, [("変化量 Δb̂_{t}", True), "（前の時刻に計算済み）"], MB_F, MB_L, lw=2.5)
box(5.95, Y2, 2.4, 0.9, [("予測 bbox", True), "b̂_{t} ＝ b_{t−1} ＋ Δb̂_{t}"], MB_F, MB_L, lw=2.5)
# 合わせて選ぶ
box(8.75, Y1 + 0.72, 1.95, 0.9, [("motion score", True), "IoU(b̂_{t}, b_{t}^{k})"], SEL_F, SEL_L)
box(11.0, Y1, 2.0, 1.15, [("score^{k} ＝", True), "0.25 × motion", "＋ 0.75 × s_{mask}^{k}"], SEL_F, SEL_L, size=14)
box(11.0, Y2 + 0.1, 2.0, 0.75, [("最大の k* を選ぶ", True)], SEL_F, SEL_L, size=14)
box(11.0, Y3, 2.0, 0.9, ["マスク k* を出力", "（SAM2 のメモリへ）"], SAM_F, SAM_L, size=13)
box(7.2, Y3, 3.3, 0.9, [("s_{mask}^{k*} ＞ 0.3 ?", True), "（SAM2 の信頼度が高いか）"], GR_F, GR_L, size=14)
box(2.4, Y4, 5.0, 1.2, [("Mamba を1ステップ進める", True), "入力：b_{t}^{k*}（画像サイズで正規化した x, y, w, h）",
                         "出力：状態 h_{t} と変化量 Δb̂_{t+1}"], MB_F, MB_L, size=14, lw=2.5)
box(8.1, Y4, 4.9, 1.2, [("Mamba を進めない", True), "次の時刻に，SAM2 の信頼度が最大の",
                         "候補で状態を作り直す"], GR_F, GR_L, size=14)
my1, my2 = Y1 + 0.45, Y2 + 0.45
arrow([(2.35, my1), (2.75, my1)]); arrow([(5.55, my1), (5.95, my1)])
arrow([(2.35, my2), (2.75, my2)]); arrow([(5.55, my2), (5.95, my2)])
arrow([(8.35, my1), (8.55, my1), (8.55, Y1 + 1.0), (8.75, Y1 + 1.0)])        # 候補 bbox → motion score
arrow([(8.35, my2), (8.55, my2), (8.55, Y1 + 1.3), (8.75, Y1 + 1.3)])        # 予測 bbox → motion score
arrow([(10.7, Y1 + 1.17), (10.85, Y1 + 1.17), (10.85, Y1 + 0.75), (11.0, Y1 + 0.75)])  # motion → score
arrow([(4.15, Y1), (4.15, Y1 - 0.2), (10.85, Y1 - 0.2), (10.85, Y1 + 0.35), (11.0, Y1 + 0.35)])  # s_mask → score
arrow([(12.0, Y1 + 1.15), (12.0, Y2 + 0.1)])                                 # score → 選ぶ
arrow([(12.0, Y2 + 0.85), (12.0, Y3)])                                       # 選ぶ → 出力
arrow([(11.0, Y2 + 0.47), (8.85, Y2 + 0.47), (8.85, Y3)])                    # 選ぶ → 判定
arrow([(7.2, Y3 + 0.45), (4.9, Y3 + 0.45), (4.9, Y4)])                       # はい
arrow([(10.5, Y3 + 0.45), (10.75, Y3 + 0.45), (10.75, Y4)])                  # いいえ
arrow([(2.4, Y4 + 0.6), (0.85, Y4 + 0.6), (0.85, Y2 + 0.9)], color=MB_L, w=2.25)  # 次の時刻へ
note(s, 0.05, Y3 + 0.05, 1.7, 0.6, "次の時刻\nt＋1 へ", size=14, color=MB_L, bold=True, align=PP_ALIGN.CENTER)
note(s, 4.0, Y3 + 0.85, 0.85, 0.3, "はい", size=13, color=TXT, align=PP_ALIGN.RIGHT)
note(s, 9.75, Y3 + 0.85, 1.0, 0.3, "いいえ", size=13, color=TXT, align=PP_ALIGN.RIGHT)
note(s, 0.35, 6.9, 12.7, 0.55,
     "・最初の5回の観測までは Mamba を使わず，予測 ＝ 直前の bbox。5回目で直近（最大12個）の bbox から状態を作り，以降は1観測につき1ステップ進める\n"
     "・Δb̂ は Mamba の出力を 1/50 にし，画像サイズ比 ±0.1 に制限。設定は表の Ours（state）を実行したときのもの（重み 0.25，しきい値 0.3）", size=11)

# ---- 5d) 動き予測器の各部分のテンソルの形（5枚目の Mamba ブロックの図をもとに）。形は epoch100.pth の重みの形と一致
s = new_slide("補足：動き予測器の各部分のテンソルの形")
LIN_F, LIN_L = RGBColor(0xDE, 0xEB, 0xF7), RGBColor(0x2E, 0x75, 0xB6)
SSM_F, SSM_L = RGBColor(0xE2, 0xF0, 0xD9), RGBColor(0x38, 0x8E, 0x3C)
ACT_F, ACT_L = RGBColor(0xFF, 0xF2, 0xCC), RGBColor(0xBF, 0x90, 0x00)
SHP = RGBColor(0x1F, 0x4E, 0x99)


def shp(x, y, text, w=2.2, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.3))
    tf = tb.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    rich(p, text, 13, True, SHP)


# 左：モデル全体（下から上へ）
note(s, 0.35, 1.15, 4.2, 0.35, "モデル全体（下から上へ）", size=15, color=TXT, bold=True)
box(0.5, 6.15, 2.6, 0.6, [("入力 b_{t}", True), "正規化した (x, y, w, h)"], GR_F, GR_L, size=13)
box(0.5, 5.15, 2.6, 0.6, [("input_proj", True), "Linear 4 → 256"], LIN_F, LIN_L, size=13)
blk = box(0.5, 3.55, 2.6, 1.2, [("Mamba ブロック ×3", True), "（3つとも同じ形）", "→ 右に中身"], MB_F, MB_L, size=13, lw=2.5)
box(0.5, 2.3, 2.6, 0.85, [("head", True), "LayerNorm → Linear 256→256", "→ LeakyReLU → Linear 256→4"], LIN_F, LIN_L, size=11)
box(0.5, 1.5, 2.6, 0.6, [("出力 Δb̂", True), "÷50 して bbox の変化に"], GR_F, GR_L, size=12)
for y1, y2 in ((6.15, 5.75), (5.15, 4.75), (3.55, 3.15), (2.3, 2.1)):
    arrow([(1.8, y1), (1.8, y2)])
shp(3.2, 6.3, "(B, T, 4)"); shp(3.2, 4.85, "(B, T, 256)"); shp(3.2, 3.2, "(B, T, 256)"); shp(3.2, 1.65, "(B, T, 4)")
# 右：Mamba ブロック1つの中
panel = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.6), Inches(1.15), Inches(8.5), Inches(5.75))
panel.adjustments[0] = 0.03; panel.fill.background(); panel.line.color.rgb = MB_L; panel.line.width = Pt(1.75); panel.shadow.inherit = False
etree.SubElement(panel.line._get_or_add_ln(), qn("a:prstDash")).set("val", "dash")
arrow([(3.1, 4.15), (4.6, 4.15)], color=MB_L, w=1.5)
note(s, 4.75, 1.2, 8.2, 0.35, "Mamba ブロック1つの中（d_model 256，d_inner 512，d_state 16）", size=14, color=MB_L, bold=True)
LX, RX2 = 6.0, 11.4          # 左の枝（x 側）・右の枝（z 側）の左端
box(7.4, 6.25, 3.0, 0.5, [("in_proj：Linear 256 → 1024", True)], LIN_F, LIN_L, size=13)
shp(10.5, 6.35, "入力 (B, T, 256)", w=2.5)
box(LX, 5.35, 2.4, 0.5, [("Conv1d（時間方向，k ＝ 4）", True)], LIN_F, LIN_L, size=12)
box(LX, 4.6, 2.4, 0.4, [("SiLU", True)], ACT_F, ACT_L, size=12)
box(LX, 3.85, 2.4, 0.45, [("x_proj：Linear 512 → 48", True)], LIN_F, LIN_L, size=12)
box(LX - 0.4, 2.75, 3.2, 0.75, [("Selective SSM", True), "状態 h：(B, 512, 16)"], SSM_F, SSM_L, size=13)
box(RX2, 4.6, 1.6, 0.4, [("SiLU", True)], ACT_F, ACT_L, size=12)
mul = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.85), Inches(2.0), Inches(0.42), Inches(0.42))
mul.fill.solid(); mul.fill.fore_color.rgb = RGBColor(0xFB, 0xE5, 0xD6); mul.line.color.rgb = SAM_L; mul.shadow.inherit = False
mp = mul.text_frame.paragraphs[0]; mp.alignment = PP_ALIGN.CENTER; rich(mp, "×", 16, True)
mul.text_frame.margin_left = mul.text_frame.margin_right = mul.text_frame.margin_top = mul.text_frame.margin_bottom = 0
box(7.55, 1.6, 2.7, 0.3, [("out_proj：Linear 512 → 256", True)], LIN_F, LIN_L, size=12)
cx, rx = LX + 1.2, RX2 + 0.8
arrow([(8.9, 6.25), (8.9, 6.05), (cx, 6.05), (cx, 5.85)])            # → x
arrow([(8.9, 6.05), (rx, 6.05), (rx, 5.0)])                         # → z
shp(cx - 1.95, 5.93, "x：(B, T, 512)", w=1.8, align=PP_ALIGN.RIGHT); shp(rx - 1.65, 5.4, "z：(B, T, 512)", w=1.6, align=PP_ALIGN.RIGHT)
arrow([(cx, 5.35), (cx, 5.0)]); arrow([(cx, 4.6), (cx, 4.3)])
shp(cx + 0.1, 4.33, "(B, T, 512)", w=1.5)
arrow([(cx, 3.85), (cx, 3.5)])
shp(cx + 0.1, 3.52, "Δ：(B, T, 16)→(B, T, 512)，B・C：(B, T, 16)", w=4.4)
arrow([(cx, 2.75), (cx, 2.21), (8.85, 2.21)])
shp(cx - 1.35, 2.3, "y：(B, T, 512)", w=1.3)
arrow([(rx, 4.6), (rx, 2.21), (9.27, 2.21)])
shp(rx - 1.6, 4.25, "(B, T, 512)", w=1.5, align=PP_ALIGN.RIGHT)
arrow([(9.06, 2.0), (9.06, 1.9)])
shp(10.35, 1.6, "出力 (B, T, 256)", w=2.6)
# 推論時に持つ状態
box(9.3, 2.55, 2.6, 0.9, [("推論で物体ごとに持つ状態", True), "conv_state：(1, 512, 4)", "ssm_state：(1, 512, 16)　×3層"], MB_F, MB_L, size=11)
note(s, 0.35, 6.95, 12.8, 0.5,
     "B：一度に処理する系列の数（学習 64，推論 1 ＝ 物体ごと）　T：系列の長さ（学習 11，推論では1観測につき 1）　パラメータ数：約138万\n"
     "層の間に残差接続や正規化はなく，3つの Mamba ブロックをそのまま重ねている。形は学習済みの重み（epoch100.pth）の形と一致", size=11)

# ---- 5e) Selective SSM の中身と，状態方程式・出力方程式の A, B, C, D の対応（1つの Mamba ブロックの中）
s = new_slide("補足：Selective SSMとA・B・C・D")
IN_F, IN_L = RGBColor(0xFB, 0xE5, 0xD6), RGBColor(0xED, 0x7D, 0x31)    # 入力から毎時刻作る
FX_F, FX_L = RGBColor(0xED, 0xED, 0xED), RGBColor(0x7F, 0x7F, 0x7F)    # 学習で決まる固定の値
box(0.35, 3.25, 2.0, 1.0, [("SSM への入力 x_{t}", True), "Conv1d → SiLU の後", "(512)"], LIN_F, LIN_L, size=12)
box(2.75, 3.25, 1.9, 1.0, [("x_proj", True), "Linear 512 → 48"], LIN_F, LIN_L, size=13)
box(5.05, 1.45, 2.4, 1.0, [("Δ_{t}（刻み幅）", True), "16 → dt_proj → softplus", "(512)"], IN_F, IN_L, size=12)
box(5.05, 3.25, 2.4, 1.0, [("B_{t}", True), "書き込み方", "(16)"], IN_F, IN_L, size=12)
box(5.05, 5.05, 2.4, 1.0, [("C_{t}", True), "読み出し方", "(16)"], IN_F, IN_L, size=12)
box(5.05, 0.0 + 6.35, 2.4, 0.6, [("D", True), "固定 (512)"], FX_F, FX_L, size=12)
box(7.85, 0.0 + 1.25 - 0.0, 2.1, 0.75, [("A ＝ −exp(A_log)", True), "固定 (512 × 16)"], FX_F, FX_L, size=12)
box(7.85, 2.3, 2.1, 0.75, [("Ā_{t} ＝ exp(Δ_{t} A)", True), "どれだけ保つか"], IN_F, IN_L, size=12)
box(7.85, 3.35, 2.1, 0.75, [("B̄_{t} ＝ Δ_{t} B_{t}", True), "どれだけ書き込むか"], IN_F, IN_L, size=12)
box(10.3, 2.05, 2.75, 1.6, [("状態方程式（記憶を更新）", True), "h_{t} ＝ Ā_{t} h_{t−1} ＋ B̄_{t} x_{t}", "状態 h：(512 × 16)"], MB_F, MB_L, size=13, lw=2.5)
box(10.3, 4.6, 2.75, 1.6, [("出力方程式（記憶を読む）", True), "y_{t} ＝ C_{t} h_{t} ＋ D x_{t}", "(512)"], MB_F, MB_L, size=13, lw=2.5)
box(10.3, 6.5, 2.75, 0.5, ["× SiLU(z) → out_proj へ"], GR_F, GR_L, size=12)
arrow([(11.67, 6.2), (11.67, 6.5)])
# 矢印
arrow([(2.35, 3.75), (2.75, 3.75)])
arrow([(4.65, 3.6), (4.85, 3.6), (4.85, 1.95), (5.05, 1.95)])     # → Δ
arrow([(4.65, 3.75), (5.05, 3.75)])                                # → B
arrow([(4.65, 3.9), (4.85, 3.9), (4.85, 5.55), (5.05, 5.55)])     # → C
arrow([(7.45, 1.95), (7.65, 1.95), (7.65, 2.55), (7.85, 2.55)])   # Δ → Ā
arrow([(8.9, 2.0), (8.9, 2.3)])                                    # A → Ā
arrow([(7.65, 2.55), (7.65, 3.6), (7.85, 3.6)])                    # Δ → B̄
arrow([(7.45, 3.75), (7.85, 3.75)])                                # B → B̄
arrow([(9.95, 2.67), (10.3, 2.67)])                                # Ā → 状態
arrow([(9.95, 3.72), (10.15, 3.72), (10.15, 3.2), (10.3, 3.2)])   # B̄ → 状態
arrow([(11.67, 3.65), (11.67, 4.6)])                               # h_t → 出力
arrow([(7.45, 5.55), (10.3, 5.55)])                                # C → 出力
arrow([(7.45, 6.65), (10.0, 6.65), (10.0, 6.05), (10.3, 6.05)])    # D → 出力
arrow([(1.35, 4.25), (1.35, 4.5), (10.6, 4.5), (10.6, 3.65)], color=LIN_L)   # x_t → 状態（B̄_t x_t）
arrow([(9.6, 4.5), (9.6, 4.85), (10.3, 4.85)], color=LIN_L)                  # x_t → 出力（D x_t）
arrow([(13.05, 2.85), (13.2, 2.85), (13.2, 1.8), (12.3, 1.8), (12.3, 2.05)], color=MB_L, w=2.0)  # h_{t-1} の引き継ぎ
shp(11.3, 1.45, "前の時刻の h_{t−1}", w=1.9)
shp(2.6, 4.53, "x_{t}（状態方程式の B̄_{t} x_{t} と，出力方程式の D x_{t} に使う）", w=6.8)
note(s, 0.35, 6.3, 4.4, 1.0, "・A は対角で，チャンネルごとに16個（512 × 16）\n・B・C は512チャンネルで共通，Δ はチャンネルごと（どれも時刻ごとに変わる）\n・3つの Mamba ブロックが，それぞれ別の A・D・x_proj を持つ", size=11)
# 凡例
for i, (f, l, txt) in enumerate(((IN_F, IN_L, "入力から毎時刻作る（Δ・B・C）：選択メカニズム"), (FX_F, FX_L, "学習で決まる固定の値（A・D）"),
                                 (MB_F, MB_L, "状態の更新と読み出し"))):
    sw = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(5.0 + 0.42 * i), Inches(0.3), Inches(0.25))
    sw.fill.solid(); sw.fill.fore_color.rgb = f; sw.line.color.rgb = l; sw.shadow.inherit = False
    note(s, 0.75, 4.97 + 0.42 * i, 4.2, 0.3, txt, size=12, color=TXT)

# ---- 6) 手法2の予備実験（MOSE）
s = new_slide("補足：手法2の予備実験（MOSE，単一物体）", "タイトルとコンテンツ")
body(s, [(0, "SAM2.1 small に時間方向の Mamba を加えて追加学習"),
         (1, "MOSE（単一物体のマスク追跡），評価は val 200 動画・570 対象，1 seed")], top=1.2, height=1.1)
rows = [["条件", "内容", "J&F↑"],
        ["P0", "SAM2.1 small（追加学習なし）", "70.08"],
        ["F0", "SAM2 全体を追加学習", "73.92"],
        ["F1", "F0 ＋ 時間方向の Mamba", "74.20"],
        ["F1-reset", "F1 で Mamba の状態を毎フレーム初期化", "74.27"]]
table(s, rows, 0.8, 2.55, 11.7, col_w=[2.0, 7.2, 2.5], size=18)
note(s, 0.8, 5.0, 11.7, 1.9,
     "F1 − F0 ＝ +0.28（95% 信頼区間 −0.67 〜 +1.16）：0 を含み，Mamba を加えた効果ははっきりしない\n"
     "F1 − F1-reset ＝ −0.06（−0.16 〜 +0.00）：状態を引き継ぐ効果も確認できない\n"
     "F0 − P0 ＝ +3.84（2.16 〜 5.73）：追加学習そのものの効果は大きい\n"
     "→ MOT での評価と，特徴・挿入位置の検討はこれから", size=16, color=TXT)

# ---- 7) checkpoint の比較
s = new_slide("補足：学習済みモデル（checkpoint）の比較", "タイトルとコンテンツ")
body(s, [(0, "表の Ours（state）は学習の最後（100 epoch）のモデル"),
         (1, "Mamba 単体の検証で最もよかったモデル（20 epoch）を SAM2 に入れると低かった")], top=1.2, height=1.1)
rows = [["モデル", "HOTA↑", "DetA↑", "AssA↑", "IDF1↑", "IDSW↓"],
        ["100 epoch（表の Ours）", "55.52", "49.60", "62.48", "64.15", "1535"],
        ["20 epoch（単体の検証で最良）", "54.61", "48.61", "61.69", "62.81", "1525"],
        ["参考：SAMURAI", "54.11", "47.89", "61.48", "62.33", "1517"]]
table(s, rows, 0.6, 2.6, 12.1, col_w=[4.1, 1.6, 1.6, 1.6, 1.6, 1.6], size=18)
note(s, 0.6, 4.7, 12.1, 0.9, "単体での良さと，SAM2 に組み込んだときの良さは一致しない。2 点の比較なので，最適な選び方はまだ決めていない", size=16, color=TXT)

_lst = prs.slides._sldIdLst
_ids = {sl.shapes.title.text_frame.text: el for sl, el in zip(prs.slides, list(_lst))}
_el = _ids["補足：動き予測器の各部分のテンソルの形"]; _lst.remove(_el); _ids["補足：動き予測器のMambaブロック"].addnext(_el)
_el2 = _ids["補足：Selective SSMとA・B・C・D"]; _lst.remove(_el2); _el.addnext(_el2)
prs.save(str(OUT))
print("saved", OUT, len(prs.slides), "slides; per-seq: wins", win, "range", round(diff[-1][0], 2), round(diff[0][0], 2))
