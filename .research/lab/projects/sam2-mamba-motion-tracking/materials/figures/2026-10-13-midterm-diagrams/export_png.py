#!/usr/bin/env python3
"""Export each slide of a pptx to PNG for quick review.

usage: export_png.py [deck.pptx] [out_dir]
  default deck: 2026-10-13-midterm-diagrams.pptx (this folder), default out: ./png
  LibreOffice is taken from $SOFFICE, else `soffice` on PATH.

The PNGs are rendered by LibreOffice. Fonts missing on this machine are substituted
(e.g. the theme font "Noto Sans JP" -> "Noto Sans CJK JP"), so PowerPoint can differ slightly.
"""

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).parent
NAMES = {  # slide title -> file name
    "従来のMOT手法": "01-tbd-tba",
    "関連研究：SAM2ベースの追跡": "02-related-work",
    "既存手法の整理：動き予測": "03-method-table",
    "研究の目的": "04-research-goal",
    "提案手法": "08-proposed-method",
    "Mambaによる時系列情報の取り込み方": "06-sliding-window-vs-state-carry",
    "Mambaを用いた動き予測": "07-motion-prediction",
    "（旧）bboxベースとマスクベースの追跡の合流": "05-old-bbox-mask-merge",
    "（旧）提案手法：IDごとのマスクを重ねた図": "09-proposed-method-stacked",
    "提案手法の課題点": "10-proposed-method-issue",
}


def titles(deck):
    try:
        from pptx import Presentation
    except ImportError:
        return []
    return [s.shapes.title.text_frame.text if s.shapes.title is not None else ""
            for s in Presentation(str(deck)).slides if s._element.get("show") != "0"]


def main():
    deck = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "2026-10-13-midterm-diagrams.pptx"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "png"
    soffice = os.environ.get("SOFFICE") or shutil.which("soffice")
    if not soffice:
        sys.exit("LibreOffice (soffice) not found: set SOFFICE=/path/to/soffice")
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        env = dict(os.environ, HOME=tmp)  # 使い捨てのプロファイルで実行する
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(deck)],
                       check=True, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pdf = Path(tmp) / (deck.stem + ".pdf")
        subprocess.run(["pdftoppm", "-png", "-r", "200", str(pdf), str(Path(tmp) / "p")], check=True)
        pages = sorted(Path(tmp).glob("p-*.png"))
        names = titles(deck)
        for i, page in enumerate(pages):
            title = names[i] if i < len(names) else ""
            name = NAMES.get(title, f"{i + 1:02d}-slide")
            shutil.copy(page, out / f"{name}.png")
            print(out / f"{name}.png")


if __name__ == "__main__":
    main()
